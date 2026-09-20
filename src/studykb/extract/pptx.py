"""PowerPoint -> Units, one per slide.

Speaker notes are included: in this corpus they often carry the sentence the
slide only gestures at.
"""

from __future__ import annotations

from pathlib import Path

from ..types import Unit


def extract(path: Path) -> list[Unit]:
    from pptx import Presentation

    prs = Presentation(str(path))
    units: list[Unit] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        title = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if not title and shape == slide.shapes.title:
                title = text.splitlines()[0]
            parts.append(text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"\n[speaker notes]\n{notes}")

        units.append(
            Unit(
                locator=f"slide {i}",
                text="\n\n".join(parts).strip(),
                provenance="text-layer",
                heading=title,
                page_no=i,
            )
        )
    return units
